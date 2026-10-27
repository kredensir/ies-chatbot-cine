"""Genera el documento del trabajo en Markdown, DOCX y DOC (RTF).

Uso:
    python scripts/make_document.py

Salidas:
    docs/Documento_IA_CineBot.md
    docs/Documento_IA_CineBot.docx
    docs/Documento_IA_CineBot.doc   (RTF compatible con Microsoft Word)
"""

from __future__ import annotations

import json
import re
import sqlite3
import struct
import sys
from datetime import date
from pathlib import Path
from typing import Any, Iterable

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from cinebot import config  # noqa: E402

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Falta python-docx. Ejecute: python -m pip install python-docx") from exc


DOCS_DIR = BASE_DIR / "docs"
MD_PATH = DOCS_DIR / "Documento_IA_CineBot.md"
DOCX_PATH = DOCS_DIR / "Documento_IA_CineBot.docx"
DOC_PATH = DOCS_DIR / "Documento_IA_CineBot.doc"

PRIMARY = RGBColor(0x6C, 0x3D, 0xFF)
DARK = RGBColor(0x1A, 0x1D, 0x2B)
MUTED = RGBColor(0x5E, 0x6A, 0x7A)
CORAL = RGBColor(0xC0, 0x39, 0x4A)

Block = tuple[Any, ...]


# ---------------------------------------------------------------------------
# Datos dinámicos
# ---------------------------------------------------------------------------
def load_metrics() -> dict[str, Any]:
    path = config.model_path("metadata")
    if not path.exists():
        raise SystemExit("No existe models/metadata.json. Ejecute scripts/train_models.py")
    return json.loads(path.read_text(encoding="utf-8"))


def load_db_stats() -> dict[str, int]:
    stats = {
        "items": 0,
        "movies": 0,
        "series": 0,
        "users": 0,
        "demo_users": 0,
        "ratings": 0,
        "app_ratings": 0,
        "messages": 0,
        "recommendations": 0,
        "genres": 0,
    }
    if not config.DB_PATH.exists():
        return stats
    with sqlite3.connect(config.DB_PATH) as connection:
        def one(sql: str) -> int:
            row = connection.execute(sql).fetchone()
            return int(row[0]) if row and row[0] is not None else 0

        stats["items"] = one("SELECT COUNT(*) FROM items")
        stats["movies"] = one("SELECT COUNT(*) FROM items WHERE media_type = 'movie'")
        stats["series"] = one("SELECT COUNT(*) FROM items WHERE media_type = 'series'")
        stats["users"] = one("SELECT COUNT(*) FROM users")
        stats["demo_users"] = one("SELECT COUNT(*) FROM users WHERE is_demo = 1")
        stats["ratings"] = one("SELECT COUNT(*) FROM ratings WHERE source = 'movielens'")
        stats["app_ratings"] = one("SELECT COUNT(*) FROM ratings WHERE source = 'app'")
        stats["messages"] = one("SELECT COUNT(*) FROM messages")
        stats["recommendations"] = one("SELECT COUNT(*) FROM recommendations")
        stats["genres"] = one("SELECT COUNT(*) FROM genres")
    return stats


# ---------------------------------------------------------------------------
# Bloques de contenido
# ---------------------------------------------------------------------------
def build_blocks(metrics: dict[str, Any], db: dict[str, int]) -> list[Block]:
    collab = metrics.get("collaborative", {})
    baseline = metrics.get("baseline", {})
    ranking = metrics.get("ranking", {})
    hybrid = ranking.get("hybrid", {})
    popularity = ranking.get("popularidad", {})
    content_rank = ranking.get("content", {})
    cf_rank = ranking.get("collaborative", {})
    intent = metrics.get("intent", {})
    clusters = metrics.get("clusters", {})
    weights = metrics.get("hybrid_weights", {})
    splits = metrics.get("splits", {})
    catalog = metrics.get("catalog", {})
    svd = metrics.get("svd_params", {})
    cluster_names = metrics.get("clusters_names", {})
    cluster_desc = metrics.get("clusters_descriptions", {})

    improvement = 0.0
    if baseline.get("rmse"):
        improvement = (baseline["rmse"] - collab.get("rmse", 0.0)) / baseline["rmse"] * 100

    cluster_rows = []
    for key in sorted(cluster_names, key=lambda value: int(value)):
        cluster_rows.append(
            [
                f"Grupo {key}",
                str(cluster_names[key]),
                str(cluster_desc.get(key, "")).replace("Géneros afines: ", ""),
            ]
        )

    blocks: list[Block] = [
        ("cover_title", "CineBot"),
        (
            "cover_subtitle",
            "Chatbot inteligente de recomendación de películas y series con Python, Flet, SQLite y aprendizaje automático",
        ),
        ("cover_rule", ""),
        ("cover_meta", "Trabajo práctico integrador"),
        ("cover_meta", "Programación · Bases de Datos · Inteligencia Artificial"),
        ("cover_meta", "Autor/a: ______________________________________"),
        ("cover_meta", "Docente: ______________________________________"),
        ("cover_meta", f"Fecha: {date.today().strftime('%d/%m/%Y')}"),
        ("cover_meta", "Carpeta del proyecto: ia/"),
        ("pagebreak",),

        ("h1", "Resumen ejecutivo"),
        (
            "p",
            "CineBot es una aplicación de escritorio y web construida con Python y Flet que recomienda "
            "películas y series mediante lenguaje natural. El sistema integra tres áreas del trabajo: "
            "**programación** orientada a objetos y por capas, **bases de datos** relacionales con SQLite, "
            "y **machine learning** clásico con modelos entrenados localmente. No utiliza un LLM externo ni "
            "depende de claves de API: todas las respuestas se generan a partir de datos, reglas y modelos "
            "que se ejecutan en el equipo del usuario.",
        ),
        (
            "p",
            "El catálogo reúne 9.832 títulos (9.742 películas y 90 series). Las películas y las 100.836 "
            "valoraciones históricas provienen de MovieLens latest-small; las series se enriquecieron con "
            "metadatos públicos de TVmaze y sinopsis originales en español. Sobre esos datos se entrenaron "
            "un modelo de contenido TF-IDF, un modelo colaborativo FunkSVD con sesgos, un recomendador "
            "híbrido con pesos ajustados por validación, una segmentación de usuarios con K-Means y un "
            "clasificador de intenciones con regresión logística.",
        ),
        (
            "p",
            f"El recomendador colaborativo alcanzó un RMSE de **{collab.get('rmse', 0):.4f}** frente a "
            f"**{baseline.get('rmse', 0):.4f}** de la media global, una mejora del **{improvement:.2f} %**. "
            f"El modelo híbrido obtuvo un NDCG@10 de **{hybrid.get('ndcg@10', 0):.4f}**, superando al "
            f"contenido puro ({content_rank.get('ndcg@10', 0):.4f}) y al colaborativo puro "
            f"({cf_rank.get('ndcg@10', 0):.4f}). El clasificador de intenciones logró una accuracy de "
            f"**{intent.get('accuracy', 0):.2%}** sobre un corpus sintético de 18 intenciones, y el "
            f"clustering de usuarios eligió k = **{clusters.get('best_k', 0)}** con una silueta de "
            f"**{clusters.get('silhouette', 0):.4f}**.",
        ),
        (
            "p",
            "El resultado es una aplicación funcional, reproducible y explicable: cada recomendación muestra "
            "sus motivos, cada conversación queda registrada en la base de datos y cada métrica se puede "
            "regenerar con un comando. El documento describe las decisiones de diseño, los fundamentos "
            "conceptuales, la implementación, los experimentos y las conclusiones del trabajo.",
        ),

        ("h1", "1. Introducción"),
        ("h2", "1.1. Contexto y motivación"),
        (
            "p",
            "La oferta de cine y series creció hasta volverse difícil de explorar. Los sistemas de "
            "recomendación reducen esa sobrecarga al estimar qué contenido puede interesar a cada persona. "
            "Detrás de una recomendación aparentemente simple conviven problemas de representación de datos, "
            "recuperación de información, predicción de preferencias, evaluación experimental e interacción "
            "con el usuario.",
        ),
        (
            "p",
            "CineBot propone una solución completa y de escala académica: un chatbot que entiende pedidos en "
            "español, consulta un catálogo heterogéneo, combina señales de contenido y de comportamiento, y "
            "presenta los resultados en una interfaz gráfica. El objetivo no es competir con un servicio "
            "industrial, sino demostrar de forma integrada los conceptos de las materias mediante una "
            "aplicación real, ejecutable y medible.",
        ),
        ("h2", "1.2. Problema abordado"),
        ("bullets", [
            "Un usuario nuevo no tiene historial: el sistema debe resolver el **arranque en frío**.",
            "Un usuario con historial necesita recomendaciones personalizadas y no repetitivas.",
            "El catálogo mezcla películas y series con metadatos incompletos y de distinta naturaleza.",
            "El lenguaje natural es ambiguo: “algo para reír” combina un ánimo con un género.",
            "Toda recomendación debe poder explicarse y quedar registrada para auditoría.",
        ]),
        ("h2", "1.3. Objetivos"),
        ("p", "**Objetivo general:** construir una aplicación de chatbot que recomiende películas y series, aplicando conceptos de programación, bases de datos y machine learning sobre datos reales."),
        ("p", "**Objetivos específicos:**"),
        ("bullets", [
            "Integrar dos fuentes de datos reales en un catálogo unificado y reproducible.",
            "Diseñar una base SQLite normalizada con integridad, vistas, triggers y búsqueda full-text.",
            "Entrenar y evaluar un recomendador híbrido de contenido y filtrado colaborativo.",
            "Implementar un clasificador de intenciones en español y un motor de diálogo con estado.",
            "Desarrollar una interfaz Flet funcional, responsive y con explicaciones visibles.",
            "Documentar el proceso, las métricas y las limitaciones de forma verificable.",
        ]),
        ("h2", "1.4. Alcance"),
        (
            "p",
            "El sistema funciona de forma local y monousuario para el perfil de invitado, con perfiles de "
            "demostración de MovieLens para probar el filtrado colaborativo. No incluye autenticación, "
            "recomendación en tiempo real a gran escala ni sinopsis de películas. La arquitectura, sin "
            "embargo, deja preparadas las tablas y los puntos de extensión para una versión multiusuario.",
        ),

        ("h1", "2. Marco conceptual"),
        ("h2", "2.1. Programación"),
        (
            "p",
            "La aplicación se organiza en capas: presentación (Flet), diálogo (ChatEngine), comprensión del "
            "lenguaje (NLU), dominio, acceso a datos y machine learning. Las dependencias apuntan hacia el "
            "dominio, no al revés. Esto permite probar el motor por consola sin abrir una ventana y cambiar "
            "la interfaz sin reentrenar modelos.",
        ),
        ("bullets", [
            "**POO y dataclasses:** `MediaItem`, `UserProfile`, `QueryFilters`, `NLUResult`, `Recommendation` y `ChatResponse` forman un contrato explícito entre capas.",
            "**Módulos y paquetes:** `cinebot.data`, `cinebot.db`, `cinebot.ml`, `cinebot.chatbot` y `cinebot.ui` separan responsabilidades.",
            "**Type hints y `from __future__ import annotations`:** documentan las firmas y facilitan el análisis estático.",
            "**Manejo de errores:** el motor captura fallos de una consulta y devuelve una respuesta comprensible en lugar de cerrar la aplicación.",
            "**Concurrencia:** Flet ejecuta el procesamiento del mensaje en un hilo aparte para no bloquear la interfaz.",
            "**Persistencia:** repositorios especializados encapsulan SQL y evitan dispersar consultas por la UI.",
        ]),
        ("h2", "2.2. Bases de datos"),
        (
            "p",
            "Se usa SQLite por ser un motor relacional completo, transaccional y sin servidor, ideal para un "
            "trabajo reproducible. El esquema aplica el modelo relacional, la normalización hasta tercera "
            "forma normal para el catálogo, claves primarias y foráneas, restricciones de dominio, índices y "
            "transacciones ACID. Además se usan objetos avanzados: vistas, triggers y una tabla virtual FTS5.",
        ),
        ("bullets", [
            "**Normalización:** los géneros no se guardan como texto repetido; se almacenan en `genres` y se relacionan mediante `item_genres`.",
            "**Integridad referencial:** `ON DELETE CASCADE` evita registros huérfanos y `CHECK` limita tipos, notas y eventos.",
            "**Vistas:** encapsulan consultas frecuentes y simplifican la UI (`v_user_stats`, `v_user_genre_taste`, `v_top_rated`).",
            "**Triggers:** recalculan los agregados locales de un ítem cuando el usuario califica o borra una nota.",
            "**FTS5:** permite búsquedas por texto completo con tokenización Unicode y eliminación de diacríticos.",
            "**Trazabilidad:** conversaciones, mensajes, recomendaciones e interacciones quedan persistidos.",
        ]),
        ("h2", "2.3. Machine learning"),
        (
            "p",
            "Se combinan técnicas supervisadas y no supervisadas. El filtrado colaborativo y el clasificador "
            "de intenciones son supervisados; K-Means es no supervisado. El recomendador de contenido es una "
            "técnica de recuperación de información basada en representaciones vectoriales. La evaluación "
            "distingue predicción de notas (regresión) de ordenamiento de recomendaciones (ranking).",
        ),
        ("bullets", [
            "**TF-IDF:** convierte texto en vectores ponderados por frecuencia e inversa de frecuencia documental.",
            "**Similitud coseno:** mide el ángulo entre vectores y es apropiada para textos de longitud variable.",
            "**FunkSVD:** factorización de matrices con sesgos y SGD, orientada a predecir notas.",
            "**K-Means:** agrupa usuarios con perfiles de género parecidos.",
            "**Regresión logística:** clasifica intenciones a partir de n-gramas de palabras y caracteres.",
            "**Validación:** partición train/validación/test, validación cruzada y early stopping para controlar el sobreajuste.",
        ]),
        ("h2", "2.4. Procesamiento de lenguaje natural"),
        (
            "p",
            "El NLU separa dos problemas. Primero, un clasificador aprende patrones generales de intención. "
            "Segundo, un conjunto de reglas y diccionarios extrae entidades concretas: géneros, ánimos, años, "
            "títulos, notas y cantidades. Esta división es deliberada: el modelo generaliza y las reglas "
            "garantizan precisión sobre el catálogo. Un paso final de refinamiento corrige la intención cuando "
            "hay evidencia léxica fuerte, por ejemplo “de qué trata X” o “califica X con 4”.",
        ),

        ("h1", "3. Arquitectura del sistema"),
        (
            "p",
            "La arquitectura sigue un flujo descendente desde la interfaz hasta los datos, con dependencias "
            "explícitas y bajo acoplamiento. El siguiente esquema resume las capas y sus responsabilidades.",
        ),
        ("code", """┌──────────────────────────────────────────────────────────────┐
│                    Interfaz Flet (app.py)                     │
│  perfiles · chat · tarjetas · sugerencias · métricas          │
└───────────────────────────────┬──────────────────────────────┘
                                │ ChatResponse / eventos
┌───────────────────────────────▼──────────────────────────────┐
│               ChatEngine (motor de diálogo)                   │
│  contexto · intenciones · calificaciones · persistencia      │
└───────────────┬───────────────────────────────┬──────────────┘
                │                               │
┌───────────────▼──────────────┐ ┌──────────────▼─────────────┐
│ NLU: intención + entidades   │ │ Recomendador híbrido       │
│ TF-IDF + LogisticRegression  │ │ contenido + FunkSVD +      │
│ reglas y refinamiento        │ │ calidad + popularidad      │
└───────────────┬──────────────┘ └──────────────┬─────────────┘
                │                               │
┌───────────────▼───────────────────────────────▼─────────────┐
│               SQLite: catálogo, usuarios, ratings,           │
│               conversaciones, recomendaciones, métricas      │
└──────────────────────────────────────────────────────────────┘"""),
        ("table", {
            "headers": ["Capa", "Módulos", "Responsabilidad"],
            "rows": [
                ["Presentación", "app.py, cinebot/ui", "Ventana, chat, tarjetas, eventos y estado visual."],
                ["Diálogo", "chatbot/engine.py", "Orquestar NLU, recomendador, acciones y persistencia."],
                ["NLU", "chatbot/nlu.py, ml/intent.py", "Predecir intención y extraer entidades."],
                ["Recomendación", "ml/recommender.py", "Generar, puntuar, diversificar y explicar candidatos."],
                ["Modelos", "ml/content.py, collaborative.py, clustering.py", "Entrenamiento, inferencia y persistencia de artefactos."],
                ["Datos", "db/repository.py, db/connection.py", "Consultas parametrizadas, transacciones y mapeo al dominio."],
                ["Datos fuente", "scripts/, data/", "Descarga, limpieza, catálogo y construcción de SQLite."],
            ],
            "caption": "Tabla 1. Capas de la arquitectura y responsabilidades.",
        }),
        ("h2", "3.1. Flujo de una consulta"),
        ("numbered", [
            "El usuario escribe un mensaje en el campo de texto o pulsa una sugerencia.",
            "La interfaz lo envía al `ChatEngine` en un hilo de trabajo.",
            "El NLU normaliza el texto, predice la intención y extrae entidades y filtros.",
            "El motor decide la acción: recomendar, buscar, informar, calificar, listar historial o resetear.",
            "El recomendador consulta candidatos en SQLite, puntúa con los modelos y aplica diversidad.",
            "El motor redacta una respuesta, la guarda en `messages` y registra las recomendaciones.",
            "La interfaz reemplaza el indicador de escritura por el mensaje, las tarjetas y nuevas sugerencias.",
        ]),
        ("h2", "3.2. Decisiones de diseño"),
        ("bullets", [
            "El motor no importa Flet: puede probarse por consola y reutilizarse en web o escritorio.",
            "Los artefactos se serializan con `joblib`; la base y los modelos se versionan mediante `metadata.json`.",
            "La base incluida permite abrir la app sin descargas; los scripts permiten reconstruirla.",
            "Las series usan IDs propios (900000 + tvmaze_id) para convivir con los IDs de MovieLens.",
            "Las notas de la app se separan de las de MovieLens mediante la columna `source`, sin sobrescribir el histórico.",
        ]),

        ("h1", "4. Datos y preparación"),
        ("h2", "4.1. Fuentes"),
        ("table", {
            "headers": ["Fuente", "Contenido", "Uso en el proyecto"],
            "rows": [
                ["MovieLens latest-small", f"{db['ratings']:,} valoraciones, {db['movies']:,} películas, 610 usuarios", "Entrenamiento y evaluación del CF; perfiles de demostración."],
                ["TVmaze API", "Metadatos de 90 series seleccionadas", "Catálogo de series, temporadas, episodios y géneros."],
                ["Sinopsis originales", "90 reseñas breves en español", "Documentos del modelo de contenido y fichas de series."],
                ["MovieLens links", "tmdbId e imdbId", "URL de ficha externa de cada película."],
            ],
            "caption": "Tabla 2. Origen y uso de los datos.",
        }),
        (
            "p",
            "MovieLens es un conjunto académico ampliamente utilizado. Su versión latest-small contiene "
            "valoraciones reales con timestamps y permite experimentar con filtrado colaborativo sin "
            "infraestructura pesada. TVmaze aporta metadatos de series, un tipo de contenido que MovieLens no "
            "cubre. La combinación obliga a resolver identificadores, géneros y campos faltantes de forma "
            "explícita.",
        ),
        ("h2", "4.2. Limpieza y unificación"),
        ("bullets", [
            "Normalización de títulos: minúsculas, sin tildes, sin año y reubicación de artículos finales (“Godfather, The” → “godfather”).",
            "Parseo del año embebido en el título de MovieLens y separación de géneros por el separador `|`.",
            "Normalización de series: `MINDHUNTER` → `Mindhunter`, `Stargate SG·1` → `Stargate SG-1`, `Shōgun` → `Shogun`.",
            "Asignación de IDs de series en el rango 900000 para evitar colisiones.",
            "Eliminación de cadenas `nan` y campos vacíos en sinopsis, URLs y pósters.",
            "Cálculo de la nota bayesiana y de la popularidad para desempatar títulos con pocos votos.",
            "Construcción de un único `catalog.csv` con 9.832 filas y columnas homogéneas.",
        ]),
        ("table", {
            "headers": ["Indicador", "Valor"],
            "rows": [
                ["Ítems del catálogo", f"{catalog.get('items', db['items']):,}"],
                ["Películas", f"{catalog.get('movies', db['movies']):,}"],
                ["Series", f"{catalog.get('series', db['series']):,}"],
                ["Valoraciones MovieLens", f"{catalog.get('movielens_ratings', db['ratings']):,}"],
                ["Usuarios históricos", f"{catalog.get('users', 610):,}"],
                ["Géneros distintos", f"{db['genres']:,}"],
                ["Valoraciones creadas en la app", f"{db['app_ratings']:,}"],
            ],
            "caption": "Tabla 3. Estadísticas del catálogo y de la base incluida.",
        }),
        ("figure", {"path": "reports/figures/catalog_composition.png", "caption": "Figura 1. Composición del catálogo: películas y series.", "width": 5.0}),
        ("figure", {"path": "reports/figures/genre_distribution.png", "caption": "Figura 2. Géneros con más títulos en el catálogo.", "width": 6.0}),
        ("figure", {"path": "reports/figures/ratings_distribution.png", "caption": "Figura 3. Distribución de las valoraciones de MovieLens.", "width": 6.0}),

        ("h1", "5. Diseño de la base de datos"),
        (
            "p",
            "El esquema se implementó en SQLite 3 y se versiona en `database/schema.sql`. La base "
            "`database/cinebot.db` se genera con un script idempotente que crea el esquema, importa el "
            "catálogo y carga las valoraciones. El diseño busca integridad, consultas eficientes y "
            "trazabilidad del comportamiento del chatbot.",
        ),
        ("h2", "5.1. Tablas principales"),
        ("table", {
            "headers": ["Tabla", "Función", "Claves y restricciones"],
            "rows": [
                ["users", "Perfiles locales y de MovieLens", "PK user_id; source e is_demo con CHECK."],
                ["items", "Catálogo unificado de películas y series", "PK item_id; media_type con CHECK; índices por año, tipo y nota."],
                ["genres", "Catálogo de géneros", "PK genre_id; nombre único."],
                ["item_genres", "Relación muchos a muchos ítem-género", "PK compuesta; FKs con ON DELETE CASCADE."],
                ["ratings", "Valoraciones históricas y de la app", "PK (user_id, item_id, source); rating entre 0,5 y 5."],
                ["feedback", "Like, dislike, visto, no interesado", "UNIQUE por usuario, ítem y acción."],
                ["interactions", "Eventos analíticos", "Eventos validados con CHECK."],
                ["conversations", "Sesiones de chat", "FK a users; fecha de inicio y fin."],
                ["messages", "Mensajes de usuario y asistente", "FK a conversations; intención, confianza y entidades JSON."],
                ["recommendations", "Recomendaciones mostradas", "FK a message, user e item; rank, score y motivo."],
                ["training_runs", "Ejecuciones de entrenamiento", "Parámetros y totales del run."],
                ["model_metrics", "Métricas por modelo y run", "UNIQUE (model, metric, run)."],
            ],
            "caption": "Tabla 4. Tablas del esquema relacional.",
        }),
        ("h2", "5.2. Relaciones y normalización"),
        (
            "p",
            "La relación entre ítems y géneros es muchos a muchos y se resuelve con `item_genres`. Las "
            "valoraciones relacionan usuarios e ítems y añaden `source` a la clave para que una nota de la app "
            "pueda convivir con la nota histórica de MovieLens sin violar la unicidad. Las conversaciones "
            "agrupan mensajes y cada mensaje puede generar varias recomendaciones, lo que produce una "
            "jerarquía uno-a-muchos y muchos-a-uno bien definida.",
        ),
        ("code", """-- Integridad y trazabilidad de una recomendación
CREATE TABLE recommendations (
    recommendation_id INTEGER PRIMARY KEY AUTOINCREMENT,
    message_id        INTEGER NOT NULL,
    user_id           INTEGER NOT NULL,
    item_id           INTEGER NOT NULL,
    rank              INTEGER NOT NULL,
    score             REAL    NOT NULL,
    reason            TEXT    NOT NULL DEFAULT '',
    FOREIGN KEY (message_id) REFERENCES messages(message_id) ON DELETE CASCADE,
    FOREIGN KEY (user_id)    REFERENCES users(user_id)    ON DELETE CASCADE,
    FOREIGN KEY (item_id)    REFERENCES items(item_id)    ON DELETE CASCADE
);"""),
        ("h2", "5.3. Triggers, vistas y FTS5"),
        (
            "p",
            "Los triggers mantienen `local_rating_sum`, `local_rating_count` y `local_rating_avg` en `items` "
            "cada vez que se inserta, actualiza o elimina una nota de la app. Así la lectura de una tarjeta no "
            "necesita agregar toda la tabla de valoraciones. Las vistas encapsulan estadísticas y rankings. "
            "FTS5 indexa título, título normalizado, sinopsis y palabras clave con eliminación de "
            "diacríticos, lo que permite buscar “parasitos” y encontrar “Parasite”.",
        ),
        ("code", """-- Búsqueda full-text tolerante a tildes
SELECT i.item_id, i.title, i.media_type
FROM items_fts f
JOIN items i ON i.item_id = f.rowid
WHERE items_fts MATCH ?
ORDER BY bm25(items_fts)
LIMIT 10;"""),
        (
            "p",
            "La conexión usa `PRAGMA foreign_keys = ON`, `journal_mode = WAL` y transacciones explícitas para "
            "las operaciones que modifican varias tablas. Las consultas se construyen con parámetros, no con "
            "concatenación de cadenas, para evitar inyección SQL.",
        ),

        ("h1", "6. Recomendador de contenido"),
        ("h2", "6.1. Representación de los ítems"),
        (
            "p",
            "Cada película o serie se transforma en un documento textual. El título se repite tres veces para "
            "darle peso; se agregan los géneros canónicos y sus equivalentes en español, las etiquetas de "
            "ánimo derivadas del género, las palabras clave, la sinopsis, la década y el tipo de contenido. "
            "Esta representación permite que una consulta como “ciencia ficción espacial” se acerque a títulos "
            "con esos términos aunque el título no los contenga.",
        ),
        ("code", """documento = tipo + título×3 + géneros_es + ánimos
          + palabras_clave×2 + sinopsis + década"""),
        ("h2", "6.2. TF-IDF y similitud coseno"),
        (
            "p",
            "Se ajusta un `TfidfVectorizer` con n-gramas de una y dos palabras, `min_df=2`, `max_df=0,85`, "
            f"`sublinear_tf` y normalización L2. El resultado es una matriz dispersa de "
            f"**{metrics.get('content', {}).get('items', 0):,} × {metrics.get('content', {}).get('vocabulary_size', 0):,}** "
            "y un vocabulario de "
            f"**{metrics.get('content', {}).get('vocabulary_size', 0):,} términos**. Para cada ítem se "
            "precomputan sus 40 vecinos más similares procesando la matriz por bloques, evitando materializar "
            "una matriz densa de 9.832 × 9.832.",
        ),
        ("code", """tf(t,d) = 1 + log(frecuencia(t,d))
idf(t)   = log((1 + N) / (1 + df(t))) + 1
w(t,d)   = tf(t,d) · idf(t) / ||vector(d)||

sim(A,B) = (A · B) / (||A|| · ||B||)"""),
        (
            "p",
            "El perfil de contenido del usuario se construye como una suma ponderada de los vectores de los "
            "ítems que valoró. Las notas altas pesan más; las bajas aportan una señal negativa pequeña y se "
            "descartan si no superan el umbral. La consulta combina el perfil del usuario con un vector de "
            "géneros, ánimos y palabras clave, y con el ítem de referencia cuando se pide “algo parecido a”.",
        ),
        (
            "p",
            "La principal ventaja es que resuelve el arranque en frío y produce explicaciones directas. La "
            "limitación es que solo captura similitud textual: dos películas pueden compartir género y "
            "sinopsis sin gustar al mismo público. Por eso se combina con el filtrado colaborativo.",
        ),

        ("h1", "7. Filtrado colaborativo con FunkSVD"),
        ("h2", "7.1. Modelo"),
        (
            "p",
            "FunkSVD descompone la matriz de valoraciones en factores latentes de usuario y de ítem. El "
            "modelo incluye la media global y sesgos por usuario y por ítem, de modo que una película "
            "universalmente mejor valorada reciba un sesgo positivo y un usuario exigente reciba uno negativo.",
        ),
        ("code", """r̂(u,i) = μ + b_u + b_i + p_u · q_i

min Σ (r(u,i) - r̂(u,i))² + λ(||p_u||² + ||q_i||² + b_u² + b_i²)"""),
        ("p", "El aprendizaje se realiza con descenso de gradiente estocástico. Para cada valoración se actualizan el error y los parámetros:"),
        ("code", """e = r(u,i) - r̂(u,i)
b_u ← b_u + α (e - λ b_u)
b_i ← b_i + α (e - λ b_i)
p_u ← p_u + α (e q_i - λ p_u)
q_i ← q_i + α (e p_u - λ q_i)"""),
        ("table", {
            "headers": ["Hiperparámetro", "Valor"],
            "rows": [
                ["Factores latentes", str(svd.get("n_factors", 24))],
                ["Épocas", str(svd.get("epochs", 35))],
                ["Tasa de aprendizaje", f"{svd.get('learning_rate', 0.008):.3f}"],
                ["Regularización L2", f"{svd.get('reg', 0.08):.2f}"],
                ["Paciencia (early stopping)", str(svd.get("patience", 6))],
                ["Usuarios en train", f"{svd.get('n_users', 0):,}"],
                ["Ítems en train", f"{svd.get('n_items', 0):,}"],
                ["Mejor época", str(svd.get("best_epoch", 0))],
                ["Mejor RMSE de validación", f"{svd.get('best_val_rmse', 0):.4f}"],
            ],
            "caption": "Tabla 5. Configuración del modelo FunkSVD.",
        }),
        ("h2", "7.2. Partición y evaluación"),
        (
            "p",
            f"Las valoraciones se dividen por usuario en entrenamiento ({splits.get('train', 0):,}), "
            f"validación ({splits.get('validation', 0):,}) y prueba ({splits.get('test', 0):,}). La partición "
            "por usuario garantiza que todos los perfiles aporten a las tres partes y evita que un usuario "
            "aparezca solo en test. El modelo se evalúa con RMSE, MAE y R²; la media global se usa como "
            "baseline.",
        ),
        ("table", {
            "headers": ["Métrica", "Resultado"],
            "rows": [
                ["RMSE FunkSVD", f"{collab.get('rmse', 0):.4f}"],
                ["MAE FunkSVD", f"{collab.get('mae', 0):.4f}"],
                ["R² FunkSVD", f"{collab.get('r2', 0):.4f}"],
                ["RMSE baseline", f"{baseline.get('rmse', 0):.4f}"],
                ["Mejora vs baseline", f"{improvement:.2f} %"],
                ["Valoraciones evaluadas", f"{int(collab.get('n', 0) or 0):,}"],
            ],
            "caption": "Tabla 6. Predicción de notas en el conjunto de test.",
        }),
        ("figure", {"path": "reports/figures/svd_learning_curve.png", "caption": "Figura 4. Curva de aprendizaje de FunkSVD: RMSE de entrenamiento y validación por época.", "width": 6.0}),
        ("h2", "7.3. Fold-in para usuarios nuevos"),
        (
            "p",
            "Un usuario nuevo no tiene fila en la matriz de factores. En lugar de reentrenar todo el modelo, "
            "se estima su vector resolviendo una regresión ridge sobre los factores de los ítems que ya "
            "valoró:",
        ),
        ("code", """p_u = (Qᵀ Q + λI)⁻¹ Qᵀ (r - μ - b_u - b_i)"""),
        (
            "p",
            "Si el usuario tiene al menos dos valoraciones, la predicción mezcla el factor histórico con el "
            "factor fold-in. Esto permite que el perfil de invitado mejore a medida que califica títulos sin "
            "reentrenar los 610 perfiles de MovieLens.",
        ),

        ("h1", "8. Recomendador híbrido"),
        ("h2", "8.1. Función de puntuación"),
        (
            "p",
            "El puntaje final combina cuatro señales normalizadas entre 0 y 1. La calidad usa la nota "
            "bayesiana de las películas y la nota media de las series; la popularidad se normaliza por el "
            "máximo del catálogo. Los pesos se aprendieron con una búsqueda en grilla sobre el conjunto de "
            "validación, maximizando NDCG@10.",
        ),
        ("code", """score(u,i) = w_c · contenido(u,i)
           + w_f · colaborativo(u,i)
           + w_q · calidad(i)
           + w_p · popularidad(i)

sujeto a  w_c + w_f + w_q + w_p = 1"""),
        ("table", {
            "headers": ["Peso", "Valor aprendido"],
            "rows": [
                ["Contenido", f"{weights.get('content', 0):.2f}"],
                ["Colaborativo", f"{weights.get('collaborative', 0):.2f}"],
                ["Calidad", f"{weights.get('quality', 0):.2f}"],
                ["Popularidad", f"{weights.get('popularity', 0):.2f}"],
            ],
            "caption": "Tabla 7. Pesos del recomendador híbrido elegidos por validación.",
        }),
        ("h2", "8.2. Diversidad, filtros y fallback"),
        ("bullets", [
            "**Filtros:** tipo, géneros, ánimos, rango de años, nota mínima y máximo de temporadas.",
            "**Exclusión:** no se recomiendan ítems ya vistos, rechazados o valorados.",
            "**Diversidad:** se resta una penalización proporcional al máximo solapamiento Jaccard de géneros con los ítems ya seleccionados.",
            "**Fallback progresivo:** si no hay suficientes candidatos, se relajan en orden la nota mínima, la época, el género y el tipo de contenido; la respuesta informa qué se relajó.",
            "**Series:** al no tener factores colaborativos, el peso del CF se redistribuye un 60 % a contenido y un 40 % a calidad.",
            "**Cold start:** un usuario con menos de tres valoraciones usa pesos específicos sin colaborativo.",
        ]),
        ("h2", "8.3. Resultados de ranking"),
        ("table", {
            "headers": ["Modelo", "Precision@10", "Recall@10", "NDCG@10", "Hit Rate@10", "Cobertura"],
            "rows": [
                ["Popularidad (baseline)", f"{popularity.get('precision@10', 0):.4f}", f"{popularity.get('recall@10', 0):.4f}", f"{popularity.get('ndcg@10', 0):.4f}", f"{popularity.get('hit_rate@10', 0):.4f}", f"{popularity.get('coverage', 0):.2%}"],
                ["Contenido", f"{content_rank.get('precision@10', 0):.4f}", f"{content_rank.get('recall@10', 0):.4f}", f"{content_rank.get('ndcg@10', 0):.4f}", f"{content_rank.get('hit_rate@10', 0):.4f}", f"{content_rank.get('coverage', 0):.2%}"],
                ["Colaborativo", f"{cf_rank.get('precision@10', 0):.4f}", f"{cf_rank.get('recall@10', 0):.4f}", f"{cf_rank.get('ndcg@10', 0):.4f}", f"{cf_rank.get('hit_rate@10', 0):.4f}", f"{cf_rank.get('coverage', 0):.2%}"],
                ["Híbrido", f"{hybrid.get('precision@10', 0):.4f}", f"{hybrid.get('recall@10', 0):.4f}", f"{hybrid.get('ndcg@10', 0):.4f}", f"{hybrid.get('hit_rate@10', 0):.4f}", f"{hybrid.get('coverage', 0):.2%}"],
            ],
            "caption": "Tabla 8. Métricas de ranking top-10 sobre el conjunto de test.",
        }),
        ("figure", {"path": "reports/figures/model_comparison.png", "caption": "Figura 5. Comparación de modelos de recomendación en test.", "width": 6.2}),
        (
            "p",
            "El baseline de popularidad obtiene un NDCG alto porque el conjunto de test premia títulos "
            "populares: muchos usuarios valoran bien las películas más conocidas. El híbrido no supera a ese "
            "baseline en precisión, pero sí mejora de forma clara al contenido puro y al colaborativo puro, y "
            "mantiene una cobertura y diversidad mayores que popularidad. El resultado refleja una lección "
            "habitual en sistemas de recomendación: un baseline simple puede ser muy competitivo cuando los "
            "datos son dispersos.",
        ),

        ("h1", "9. Segmentación de usuarios con K-Means"),
        ("h2", "9.1. Construcción de características"),
        (
            "p",
            "Cada usuario se representa por su nota media en cada género. Para evitar que una sola valoración "
            "en un género minoritario domine el perfil, se aplica un suavizado bayesiano con un prior de cinco "
            "valoraciones y la media global del género. Después se aplica un **doble centrado**: se resta la "
            "media del propio usuario (que captura su severidad o generosidad general) y la media global del "
            "género (que captura el nivel de valoración habitual de ese género). El resultado expresa qué "
            "géneros prefiere cada persona *por encima de lo esperado*. Las desviaciones se ponderan por "
            "log(1 + número de valoraciones) para reducir el ruido de los géneros con pocos datos y luego se "
            "estandarizan las columnas antes de aplicar K-Means.",
        ),
        ("code", """perfil(u,g) = (Σ ratings(u,g) + μ_g · m) / (n(u,g) + m),  m = 5
residuo(u,g) = perfil(u,g) − media_u(perfil) − μ_g + media_global"""),
        ("h2", "9.2. Elección de k y resultados"),
        (
            "p",
            "Se probaron valores de k entre 3 y 8. Para cada uno se calculó la inercia y el coeficiente de "
            "silueta; se eligió el k con mayor silueta. PCA en dos componentes se usa solo para visualizar, no "
            "para entrenar. Los centroides se interpretan con la media del perfil residual de los miembros, de "
            "modo que los géneros mostrados son los que el grupo prefiere por encima de su propia media y de la "
            "media del catálogo. La silueta es moderada porque los gustos cinematográficos son continuos y los "
            "grupos se solapan; su valor está en comparar alternativas de k y en producir segmentos "
            "interpretables, no en demostrar una separación perfecta.",
        ),
        ("table", {
            "headers": ["Parámetro", "Valor"],
            "rows": [
                ["k elegido", str(clusters.get("best_k", 0))],
                ["Silueta", f"{clusters.get('silhouette', 0):.4f}"],
                ["Usuarios agrupados", f"{int(clusters.get('n_users', 0) or 0):,}"],
                ["Varianza explicada por PCA", f"{clusters.get('explained_variance', 0):.2%}"],
            ],
            "caption": "Tabla 9. Resultados de la segmentación.",
        }),
        ("table", {
            "headers": ["Grupo", "Nombre interpretable", "Géneros afines y nota media"],
            "rows": cluster_rows,
            "caption": "Tabla 10. Interpretación de los clusters de usuarios.",
        }),
        ("figure", {"path": "reports/figures/user_clusters.png", "caption": "Figura 6. Segmentación de usuarios con K-Means y proyección PCA en dos dimensiones.", "width": 5.8}),
        (
            "p",
            "El perfil se muestra en la barra lateral de la aplicación y se actualiza al cambiar de usuario. "
            "Para un usuario nuevo, `predict_profile` construye su vector con las mismas reglas y lo asigna al "
            "centroide más cercano. Esta información se usa como contexto explicativo, no como filtro rígido: "
            "el recomendador sigue combinando contenido y comportamiento.",
        ),

        ("h1", "10. Comprensión del lenguaje natural"),
        ("h2", "10.1. Intenciones y entidades"),
        (
            "p",
            "El clasificador distingue 18 intenciones. Las entidades se extraen con diccionarios y expresiones "
            "regulares: géneros (incluidos sinónimos coloquiales), ánimos, años o décadas, cantidades, "
            "temporadas, notas y títulos. El matcher de títulos combina alias manuales, coincidencia exacta y "
            "difflib para tolerar errores de escritura. En la extracción de cantidades los artículos «un» y "
            "«una» no se interpretan como cantidad —«recomiéndame una película» pide el lote habitual de cinco—; "
            "para solicitar un único resultado se exige una fórmula explícita como «solo una» o «una sola».",
        ),
        ("table", {
            "headers": ["Familia", "Intenciones"],
            "rows": [
                ["Social", "saludo, despedida, agradecimiento, ayuda, afirmación, rechazo"],
                ["Recomendación", "recomendar_pelicula, recomendar_serie, recomendar_genero, recomendar_animo, similar, otra_recomendacion"],
                ["Consulta", "informacion, historial, top_valoradas"],
                ["Acción", "calificar, resetear"],
                ["Fuera de dominio", "desconocido"],
            ],
            "caption": "Tabla 11. Familias de intenciones del chatbot.",
        }),
        ("h2", "10.2. Clasificador"),
        (
            "p",
            "Se entrenó una regresión logística multinomial sobre una unión de TF-IDF de palabras (1-2 "
            "gramas) y de caracteres (3-5 gramas). Los caracteres ayudan con tildes, abreviaturas y errores de "
            "tipeo; las palabras aportan interpretabilidad. Se usó `class_weight='balanced'`, 2.000 "
            "iteraciones y validación cruzada de 5 pliegues.",
        ),
        ("table", {
            "headers": ["Métrica", "Valor"],
            "rows": [
                ["Accuracy", f"{intent.get('accuracy', 0):.4f}"],
                ["F1 macro", f"{intent.get('macro_f1', 0):.4f}"],
                ["F1 ponderado", f"{intent.get('weighted_f1', 0):.4f}"],
                ["Validación cruzada (media)", f"{intent.get('cv_accuracy_mean', 0):.4f}"],
                ["Desvío de la validación cruzada", f"{intent.get('cv_accuracy_std', 0):.4f}"],
                ["Ejemplos etiquetados", f"{int(intent.get('n_examples', 0) or 0):,}"],
                ["Intenciones", f"{int(intent.get('n_intents', 0) or 0):,}"],
                ["Características TF-IDF", f"{int(intent.get('n_features', 0) or 0):,}"],
            ],
            "caption": "Tabla 12. Evaluación del clasificador de intenciones.",
        }),
        ("figure", {"path": "reports/figures/intent_confusion.png", "caption": "Figura 7. Matriz de confusión del clasificador de intenciones.", "width": 6.2}),
        (
            "p",
            "La accuracy del clasificador es moderada porque el corpus es sintético, pequeño y contiene "
            "intenciones semánticamente cercanas (por ejemplo, recomendar una película genérica frente a "
            "recomendar por género). Esa limitación se compensa con el refinamiento por reglas: si el texto "
            "contiene un título y una fórmula de calificación, la intención se corrige a `calificar`; si "
            "contiene “de qué trata” y un título, se corrige a `informacion`. Así, el sistema final es más "
            "robusto que el clasificador aislado.",
        ),
        ("h2", "10.3. Estado del diálogo"),
        (
            "p",
            "El motor mantiene el tipo de contenido, los últimos filtros, la última recomendación y una "
            "acción pendiente. Gracias a ese estado, “dame otra” recupera el contexto anterior y “5” completa "
            "una calificación que había quedado a la espera del título o de la nota. Las respuestas incluyen "
            "sugerencias de seguimiento para hacer visible el estado y reducir la ambigüedad.",
        ),

        ("h1", "11. Aplicación e interfaz"),
        ("h2", "11.1. Interfaz Flet"),
        (
            "p",
            "La interfaz es una ventana oscura con estética cinematográfica. La barra lateral contiene el "
            "selector de perfil, las estadísticas del usuario, los atajos de género y el panel de métricas de "
            "los modelos. El panel principal muestra el chat, las tarjetas de recomendación, las sugerencias y "
            "el compositor. En pantallas angostas la barra lateral se oculta y el chat ocupa el ancho "
            "disponible.",
        ),
        ("bullets", [
            "**Tarjetas:** póster generado con degradado e iniciales, año, géneros, duración, nota, predicción, motivo, estrellas, like/dislike, ficha y enlace externo.",
            "**Chat:** burbujas alineadas, Markdown, desplazamiento automático y un indicador animado mientras el motor procesa.",
            "**Perfiles:** invitado local y 20 perfiles de MovieLens con historial real.",
            "**Sugerencias:** chips contextuales que cambian según la última respuesta.",
            "**Acciones:** nueva conversación, ayuda, borrar valoraciones y cambio de perfil.",
            "**Métricas:** RMSE, R², NDCG@10, Precision@10, accuracy del NLU y número de clusters.",
        ]),
        ("h2", "11.2. Ejecución y experiencia de uso"),
        ("code", """python app.py          # escritorio
python app.py --web    # navegador en http://127.0.0.1:8550"""),
        (
            "p",
            "En escritorio, Flet abre una ventana nativa; en modo web, FastAPI y Uvicorn sirven la aplicación "
            "en localhost. El procesamiento del mensaje se ejecuta con `page.run_thread` para que la ventana "
            "siga respondiendo. Las tarjetas y el panel lateral se actualizan al cambiar el perfil o al "
            "registrar una nota.",
        ),
        ("table", {
            "headers": ["Acción del usuario", "Resultado"],
            "rows": [
                ["“recomiéndame una película de ciencia ficción”", "Filtra por tipo y género, recomienda y explica."],
                ["“series de misterio cortas”", "Aplica género, tipo y máximo de temporadas."],
                ["“de qué trata Breaking Bad”", "Devuelve ficha, géneros, valoración y sinopsis."],
                ["“algo parecido a Inception”", "Usa similitud de contenido y excluye el título de origen."],
                ["“califica Inception con 4”", "Guarda la nota, actualiza triggers y muestra el perfil."],
                ["“dame otra”", "Recupera filtros y excluye las recomendaciones ya mostradas."],
                ["“borra mis valoraciones”", "Elimina notas de la app y feedback, conserva MovieLens."],
            ],
            "caption": "Tabla 13. Diálogos de ejemplo y comportamiento esperado.",
        }),

        ("h1", "12. Evaluación global y discusión"),
        ("table", {
            "headers": ["Componente", "Métrica principal", "Resultado"],
            "rows": [
                ["Predicción de notas", "RMSE", f"{collab.get('rmse', 0):.4f}"],
                ["Predicción de notas", "Mejora vs baseline", f"{improvement:.2f} %"],
                ["Ranking híbrido", "NDCG@10", f"{hybrid.get('ndcg@10', 0):.4f}"],
                ["Ranking híbrido", "Hit Rate@10", f"{hybrid.get('hit_rate@10', 0):.4f}"],
                ["Ranking híbrido", "Cobertura", f"{hybrid.get('coverage', 0):.2%}"],
                ["NLU", "Accuracy", f"{intent.get('accuracy', 0):.2%}"],
                ["NLU", "F1 macro", f"{intent.get('macro_f1', 0):.4f}"],
                ["Clustering", "Silueta", f"{clusters.get('silhouette', 0):.4f}"],
            ],
            "caption": "Tabla 14. Resumen de métricas del sistema.",
        }),
        ("p", "**Fortalezas observadas:**"),
        ("bullets", [
            "El sistema es completamente local, reproducible y no depende de servicios de pago.",
            "La separación por capas permitió probar el motor, los modelos y la base de datos de forma independiente.",
            "El híbrido mejora al contenido y al colaborativo puro en ranking, y conserva explicaciones útiles.",
            "La base de datos demuestra un uso real de integridad, triggers, vistas, FTS5 y trazabilidad.",
            "La interfaz traduce概念s técnicos en una experiencia comprensible para un usuario no técnico.",
        ]),
        ("p", "**Limitaciones:**"),
        ("bullets", [
            "MovieLens latest-small tiene 610 usuarios; las métricas de ranking tienen alta varianza y no representan a toda la población.",
            "Las películas no incluyen sinopsis; la ficha informa esa ausencia en lugar de inventar contenido.",
            "El catálogo de series es curado y limitado a 90 títulos.",
            "El corpus de intenciones es sintético; ampliarlo con mensajes reales mejoraría la generalización.",
            "El baseline de popularidad es muy competitivo en NDCG; el híbrido prioriza diversidad y explicabilidad, no solo precisión.",
            "La evaluación offline no mide satisfacción real ni retención; sería necesario un estudio con usuarios.",
        ]),

        ("h1", "13. Conclusiones"),
        (
            "p",
            "CineBot demuestra que es posible integrar programación, bases de datos y machine learning en una "
            "aplicación coherente y funcional sin depender de un LLM externo. El proyecto obligó a tomar "
            "decisiones en todas las capas: normalizar datos heterogéneos, diseñar un esquema con integridad y "
            "trazabilidad, entrenar modelos con métricas adecuadas, manejar el arranque en frío y traducir los "
            "resultados a una interfaz conversacional.",
        ),
        (
            "p",
            "Los resultados son razonables para un conjunto académico: el filtrado colaborativo reduce el RMSE "
            f"un {improvement:.1f} % respecto de la media global; el híbrido supera a sus componentes puros en "
            "NDCG@10 y mantiene una cobertura del "
            f"{hybrid.get('coverage', 0):.1%}; y el chatbot resuelve consultas de recomendación, información, "
            "calificación e historial con estado persistente. La evaluación también mostró límites reales, "
            "especialmente la dispersión de los datos y la dificultad de superar un baseline de popularidad.",
        ),
        (
            "p",
            "Como trabajo integrador, el valor principal está en la trazabilidad del proceso: cada dato tiene "
            "un origen, cada modelo tiene una métrica, cada recomendación tiene una explicación y cada "
            "conversación queda registrada. Esa cadena es la que convierte una demo en un sistema defendible.",
        ),

        ("h1", "14. Trabajo futuro"),
        ("bullets", [
            "Enriquecer sinopsis de películas con TMDB o Wikipedia manteniendo el modo offline mediante caché.",
            "Incorporar embeddings semánticos (Sentence Transformers) y comparar con TF-IDF.",
            "Probar factorización con sesgos temporales y modelos neuronales (NCF).",
            "Añadir evaluación online con clics, tiempo de visualización y encuestas de satisfacción.",
            "Convertir la app en multiusuario con autenticación y API REST.",
            "Internacionalizar la interfaz y el NLU (español/inglés).",
            "Explicaciones contrafactuales: “si te gustó X, Y es una alternativa más ligera”.",
        ]),

        ("h1", "15. Mapa de conceptos aplicados"),
        ("table", {
            "headers": ["Área", "Concepto", "Evidencia en el proyecto"],
            "rows": [
                ["Programación", "POO y dataclasses", "domain.py, modelos tipados y contratos entre capas."],
                ["Programación", "Modularidad y bajo acoplamiento", "Paquetes data, db, ml, chatbot, ui."],
                ["Programación", "Manejo de errores y concurrencia", "Captura en ChatEngine y page.run_thread."],
                ["Bases de datos", "Modelo relacional y 3FN", "items, genres e item_genres."],
                ["Bases de datos", "Integridad referencial", "FOREIGN KEY, CHECK y ON DELETE CASCADE."],
                ["Bases de datos", "Transacciones", "db.transaction() en operaciones compuestas."],
                ["Bases de datos", "Vistas y triggers", "v_user_stats, trg_ratings_ai/au/ad."],
                ["Bases de datos", "Full-text search", "items_fts con FTS5 y unicode61."],
                ["Machine learning", "TF-IDF y coseno", "content.py."],
                ["Machine learning", "Filtrado colaborativo", "collaborative.py (FunkSVD)."],
                ["Machine learning", "Sistema híbrido", "recommender.py y pesos por validación."],
                ["Machine learning", "Clustering", "clustering.py (K-Means + PCA + silueta)."],
                ["Machine learning", "Clasificación", "intent.py (regresión logística)."],
                ["Machine learning", "Evaluación", "RMSE, MAE, R², Precision, Recall, NDCG, Hit Rate."],
                ["NLP", "Intenciones y entidades", "nlu.py y corpus etiquetado."],
                ["NLP", "Estado del diálogo", "ChatEngine y acciones pendientes."],
            ],
            "caption": "Tabla 15. Correspondencia entre conceptos de las materias y su implementación.",
        }),

        ("h1", "16. Reproducibilidad"),
        ("p", "Todos los pasos se ejecutan desde la carpeta `ia`:"),
        ("code", """python -m pip install -r requirements.txt

python scripts/download_data.py --series
python scripts/build_series_catalog.py
python -m cinebot.data.catalog
python scripts/build_database.py --force --rebuild-catalog
python scripts/train_models.py
python scripts/evaluate_models.py --live
python -m unittest discover -s tests -v
python scripts/make_document.py"""),
        ("p", "Los artefactos principales son:"),
        ("bullets", [
            "`database/cinebot.db`: base SQLite lista para usar.",
            "`models/*.joblib`: modelos de contenido, colaborativo, intenciones y clustering.",
            "`models/metadata.json`: métricas, hiperparámetros y pesos del híbrido.",
            "`reports/evaluacion_modelos.md`: informe automático de evaluación.",
            "`reports/figures/*.png`: figuras del documento.",
            "`docs/Documento_IA_CineBot.doc` y `.docx`: documento final.",
        ]),

        ("h1", "17. Bibliografía y fuentes"),
        ("bullets", [
            "GroupLens Research. MovieLens Latest Datasets. https://grouplens.org/datasets/movielens/",
            "TVmaze. Public API. https://www.tvmaze.com/api",
            "The Movie Database (TMDB). https://www.themoviedb.org/",
            "Koren, Y., Bell, R., & Volinsky, C. (2009). Matrix Factorization Techniques for Recommender Systems. IEEE Computer.",
            "Rendle, S. (2010). Factorization Machines. ICDM.",
            "Järvelin, K., & Kekäläinen, J. (2002). Cumulated Gain-based Evaluation of IR Techniques. ACM TOIS.",
            "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR.",
            "Flet. Documentación oficial. https://flet.dev/",
            "SQLite. SQLite Documentation: FTS5, triggers and views. https://www.sqlite.org/docs.html",
        ]),
        ("h2", "Declaración sobre el uso de datos"),
        (
            "p",
            "MovieLens se utiliza con fines académicos y de investigación, respetando las condiciones del "
            "conjunto original. Los metadatos de series provienen de la API pública de TVmaze y las sinopsis "
            "en español fueron redactadas para este proyecto. La aplicación no redistribuye imágenes ni "
            "sinopsis protegidas de TMDB: solo usa identificadores públicos para enlazar la ficha externa.",
        ),
    ]
    return blocks


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------
def render_markdown(blocks: Iterable[Block]) -> str:
    lines: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "cover_title":
            lines += [f"# {block[1]}", ""]
        elif kind == "cover_subtitle":
            lines += [f"**{block[1]}**", ""]
        elif kind == "cover_rule":
            lines += ["---", ""]
        elif kind == "cover_meta":
            lines += [block[1], ""]
        elif kind == "pagebreak":
            lines += ["---", ""]
        elif kind == "h1":
            lines += [f"## {block[1]}", ""]
        elif kind == "h2":
            lines += [f"### {block[1]}", ""]
        elif kind == "h3":
            lines += [f"#### {block[1]}", ""]
        elif kind == "p":
            lines += [block[1], ""]
        elif kind in {"bullets", "numbered"}:
            marker = "-" if kind == "bullets" else "1."
            for item in block[1]:
                lines.append(f"{marker} {item}")
            lines.append("")
        elif kind == "table":
            data = block[1]
            headers = data["headers"]
            rows = data["rows"]
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("|" + "|".join("---" for _ in headers) + "|")
            for row in rows:
                lines.append("| " + " | ".join(str(cell) for cell in row) + " |")
            lines.append("")
            if data.get("caption"):
                lines += [f"*{data['caption']}*", ""]
        elif kind == "figure":
            data = block[1]
            path = "../" + data["path"].replace("\\", "/")
            lines += [f"![{data['caption']}]({path})", "", f"*{data['caption']}*", ""]
        elif kind == "code":
            lines += ["```text", str(block[1]).rstrip(), "```", ""]
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_inline_runs(paragraph, text: str, size: float | None = None, color=None) -> None:
    parts = re.split(r"(\*\*.+?\*\*)", str(text))
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def render_docx(blocks: Iterable[Block], path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (("Heading 1", 17, PRIMARY), ("Heading 2", 13, DARK), ("Heading 3", 11.5, MUTED)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("CineBot · Documento del trabajo · ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED
    add_page_number(footer)

    for block in blocks:
        kind = block[0]
        if kind == "cover_title":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(110)
            run = paragraph.add_run(block[1])
            run.bold = True
            run.font.size = Pt(54)
            run.font.name = "Cambria"
            run.font.color.rgb = PRIMARY
        elif kind == "cover_subtitle":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(block[1])
            run.font.size = Pt(17)
            run.font.color.rgb = DARK
            paragraph.paragraph_format.space_after = Pt(22)
        elif kind == "cover_rule":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("—" * 18)
            run.font.color.rgb = PRIMARY
        elif kind == "cover_meta":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(block[1])
            run.font.size = Pt(11.5)
            run.font.color.rgb = MUTED
        elif kind == "pagebreak":
            doc.add_page_break()
        elif kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind == "h2":
            doc.add_heading(block[1], level=2)
        elif kind == "h3":
            doc.add_heading(block[1], level=3)
        elif kind == "p":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_runs(paragraph, block[1])
        elif kind == "bullets":
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Bullet")
                add_inline_runs(paragraph, item)
        elif kind == "numbered":
            for item in block[1]:
                paragraph = doc.add_paragraph(style="List Number")
                add_inline_runs(paragraph, item)
        elif kind == "table":
            data = block[1]
            headers = data["headers"]
            rows = data["rows"]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            header_cells = table.rows[0].cells
            for index, header in enumerate(headers):
                header_cells[index].text = ""
                paragraph = header_cells[index].paragraphs[0]
                run = paragraph.add_run(str(header))
                run.bold = True
                run.font.size = Pt(8.5)
                set_cell_shading(header_cells[index], "EDE7FF")
            for row in rows:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = ""
                    paragraph = cells[index].paragraphs[0]
                    run = paragraph.add_run(str(value))
                    run.font.size = Pt(8.5)
            if data.get("caption"):
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(data["caption"])
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = MUTED
        elif kind == "figure":
            data = block[1]
            image_path = BASE_DIR / data["path"]
            if image_path.exists():
                doc.add_picture(str(image_path), width=Inches(float(data.get("width", 6.0))))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(data["caption"])
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = MUTED
        elif kind == "code":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(str(block[1]).rstrip())
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            run.font.color.rgb = DARK
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F4F6FA")
            paragraph._p.get_or_add_pPr().append(shading)

    props = doc.core_properties
    props.title = "CineBot — Chatbot de recomendación de películas y series"
    props.subject = "Trabajo práctico de Programación, Bases de Datos e Inteligencia Artificial"
    props.author = "CineBot"
    props.keywords = "Python, Flet, SQLite, machine learning, recomendación, TF-IDF, FunkSVD, K-Means, NLP"
    props.comments = "Documento generado automáticamente desde los artefactos del proyecto."
    doc.save(str(path))


# ---------------------------------------------------------------------------
# DOC (RTF)
# ---------------------------------------------------------------------------
def _rtf_unicode(text: str) -> str:
    out: list[str] = []
    for char in str(text):
        code = ord(char)
        if char == "\\":
            out.append("\\\\")
        elif char == "{":
            out.append("\\{")
        elif char == "}":
            out.append("\\}")
        elif char == "\n":
            out.append("\\line ")
        elif code > 127:
            # RTF usa enteros con signo de 16 bits para \uN.
            signed = code if code <= 32767 else code - 65536
            out.append(f"\\u{signed}?")
        else:
            out.append(char)
    return "".join(out)


def _rtf_inline(text: str) -> str:
    parts = re.split(r"(\*\*.+?\*\*)", str(text))
    out: list[str] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append("\\b " + _rtf_unicode(part[2:-2]) + "\\b0 ")
        else:
            out.append(_rtf_unicode(part))
    return "".join(out)


def _png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as handle:
        header = handle.read(24)
    if len(header) < 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
        return (1000, 600)
    width, height = struct.unpack(">II", header[16:24])
    return int(width), int(height)


def _rtf_image(path: Path, width_inches: float = 6.0) -> str:
    data = path.read_bytes()
    width, height = _png_size(path)
    goal_width = int(width_inches * 1440)
    goal_height = max(1, int(goal_width * height / max(1, width)))
    hex_data = data.hex()
    wrapped = "\n".join(hex_data[i : i + 120] for i in range(0, len(hex_data), 120))
    return (
        f"{{\\pict\\pngblip\\picw{width}\\pich{height}"
        f"\\picwgoal{goal_width}\\pichgoal{goal_height}\n{wrapped}}}"
    )


def _rtf_table(headers: list[str], rows: list[list[str]]) -> str:
    columns = len(headers)
    weights = []
    for index in range(columns):
        cells = [str(headers[index])] + [str(row[index]) for row in rows]
        weights.append(max(8, min(45, max(len(cell) for cell in cells))))
    total = sum(weights)
    usable = 9360
    widths = [max(900, int(usable * weight / total)) for weight in weights]
    scale = usable / sum(widths)
    widths = [int(width * scale) for width in widths]

    def row(cells: list[str], bold: bool = False) -> str:
        prefix = "\\trowd\\trgaph108\\trleft0"
        for index, width in enumerate(widths):
            prefix += f"\\cellx{sum(widths[: index + 1])}"
        body = ""
        for cell in cells:
            content = _rtf_inline(cell)
            if bold:
                body += "\\pard\\intbl\\fs17\\b " + content + "\\b0\\cell "
            else:
                body += "\\pard\\intbl\\fs17 " + content + "\\cell "
        return prefix + "\n" + body + "\\row\n"

    output = [row(headers, bold=True)]
    for data_row in rows:
        output.append(row(data_row))
    return "".join(output) + "\\pard\\par\n"


def render_rtf(blocks: Iterable[Block], path: Path) -> None:
    out: list[str] = [
        "{\\rtf1\\ansi\\ansicpg1252\\deff0",
        "{\\fonttbl{\\f0\\fnil\\fcharset0 Calibri;}{\\f1\\fswiss\\fcharset0 Cambria;}{\\f2\\fmodern\\fcharset0 Consolas;}}",
        "{\\colortbl ;\\red26\\green29\\blue43;\\red108\\green61\\blue255;\\red94\\green106\\blue122;\\red192\\green57\\blue74;}",
        "\\viewkind4\\uc1",
    ]

    for block in blocks:
        kind = block[0]
        if kind == "cover_title":
            out.append("\\pard\\qc\\f1\\fs72\\b\\cf2 " + _rtf_unicode(block[1]) + "\\b0\\cf1\\par")
        elif kind == "cover_subtitle":
            out.append("\\pard\\qc\\fs30\\cf1 " + _rtf_unicode(block[1]) + "\\par")
        elif kind == "cover_rule":
            out.append("\\pard\\qc\\fs24\\cf2 " + _rtf_unicode("—" * 20) + "\\par")
        elif kind == "cover_meta":
            out.append("\\pard\\qc\\fs22\\cf3 " + _rtf_inline(block[1]) + "\\cf1\\par")
        elif kind == "pagebreak":
            out.append("\\page")
        elif kind == "h1":
            out.append("\\pard\\sa180\\sb240\\f1\\fs34\\b\\cf2 " + _rtf_inline(block[1]) + "\\b0\\cf1\\par")
        elif kind == "h2":
            out.append("\\pard\\sa120\\sb180\\fs26\\b\\cf1 " + _rtf_inline(block[1]) + "\\b0\\par")
        elif kind == "h3":
            out.append("\\pard\\sa100\\sb140\\fs22\\b\\cf3 " + _rtf_inline(block[1]) + "\\b0\\cf1\\par")
        elif kind == "p":
            out.append("\\pard\\qj\\sa120\\sl276\\slmult1\\fs21 " + _rtf_inline(block[1]) + "\\par")
        elif kind == "bullets":
            for item in block[1]:
                out.append("\\pard\\qj\\sa80\\sl270\\slmult1\\li420\\fi-220\\fs21 \\bullet\\tab " + _rtf_inline(item) + "\\par")
        elif kind == "numbered":
            for index, item in enumerate(block[1], start=1):
                out.append(f"\\pard\\qj\\sa80\\sl270\\slmult1\\li420\\fi-260\\fs21 {index}.\\tab " + _rtf_inline(item) + "\\par")
        elif kind == "table":
            data = block[1]
            out.append(_rtf_table(data["headers"], data["rows"]))
            if data.get("caption"):
                out.append("\\pard\\qc\\fs17\\i\\cf3 " + _rtf_inline(data["caption"]) + "\\i0\\cf1\\par")
        elif kind == "figure":
            data = block[1]
            image_path = BASE_DIR / data["path"]
            if image_path.exists():
                out.append("\\pard\\qc " + _rtf_image(image_path, float(data.get("width", 6.0))) + "\\par")
                out.append("\\pard\\qc\\fs17\\i\\cf3 " + _rtf_inline(data["caption"]) + "\\i0\\cf1\\par")
        elif kind == "code":
            out.append("\\pard\\qj\\sa120\\sl240\\slmult1\\li260\\f2\\fs17\\cf1")
            lines = str(block[1]).rstrip().split("\n")
            out.append("\\line ".join(_rtf_unicode(line) for line in lines))
            out.append("\\f0\\par")

    out.append("}")
    path.write_bytes("\n".join(out).encode("ascii", errors="replace"))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> int:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    metrics = load_metrics()
    stats = load_db_stats()
    blocks = build_blocks(metrics, stats)

    MD_PATH.write_text(render_markdown(blocks), encoding="utf-8")
    print(f"[OK] Markdown: {MD_PATH}")

    render_docx(blocks, DOCX_PATH)
    print(f"[OK] Word (.docx): {DOCX_PATH}")

    render_rtf(blocks, DOC_PATH)
    print(f"[OK] Word-compatible (.doc/RTF): {DOC_PATH}")

    words = sum(len(str(block[1]).split()) for block in blocks if block[0] in {"p", "bullets", "numbered"})
    print(f"[OK] Documento generado con {len(blocks)} bloques y ~{words:,} palabras de cuerpo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
